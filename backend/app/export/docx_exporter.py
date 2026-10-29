import io
from typing import Any

import docx
import docx.opc.constants
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from app.export.content import ResumeContentBuilder


class DocxExporter(ResumeContentBuilder):
    """
    Renders a tailored resume as a .docx using Arial typography.
    Prefers *_parsed structured sections when available.
    """

    FONT_FAMILY = "Arial"
    COLOR_PRIMARY_HEX = "003366"
    COLOR_PRIMARY_RGB = RGBColor(0, 51, 102)
    COLOR_BODY_RGB = RGBColor(30, 30, 30)
    COLOR_MUTED_RGB = RGBColor(80, 80, 80)

    # ---------------------------------------------------------------- #
    # public entry point
    # ---------------------------------------------------------------- #

    def export_tailored(
        self,
        sections: dict,
        missing_keywords: list[str] | None = None,
        rewritten_bullets: dict | None = None,
        target_role: str = "",
        match_title: bool = False,
    ) -> bytes:
        doc = docx.Document()
        self._setup_page(doc)
        self._setup_base_style(doc)
        self._setup_bullet_style(doc)

        c = sections.get("contact_parsed", {}) or {}
        # A4: the resume's own title always wins. The job posting's title is
        # only substituted when the user explicitly opts in per export —
        # otherwise the export would silently claim a role the resume doesn't.
        self._add_header(
            doc, c, target_role, self.resume_title(sections), match_title
        )

        if sections.get("summary", {}).get("raw"):
            self._heading(doc, "SUMMARY")
            # A4: no silent edits to the candidate's own words — "Aspiring
            # Full Stack Developer" used to be rewritten to "Full Stack
            # Developer" without asking. If the user wants it gone, they can
            # change their resume.
            self._add_flow_text(doc, sections["summary"]["raw"])

        # ✅ Prefer structured education_parsed
        edu_parsed = sections.get("education_parsed")
        edu_raw = sections.get("education", {}).get("raw", "")
        if edu_parsed or edu_raw:
            self._heading(doc, "EDUCATION")
            if edu_parsed:
                for entry in edu_parsed:
                    self._add_education_entry(doc, entry)
            else:
                self._add_entries(doc, edu_raw)

        if sections.get("skills", {}).get("raw"):
            self._heading(doc, "TECHNICAL SKILLS")
            self._add_skills(doc, sections["skills"]["raw"], missing_keywords)

        exp_parsed = sections.get("experience_parsed")
        exp_raw = sections.get("experience", {}).get("raw", "")
        if exp_parsed or exp_raw:
            self._heading(doc, "EXPERIENCE")
            if exp_parsed:
                for entry in exp_parsed:
                    self._add_structured_entry(doc, entry, rewritten_bullets)
            else:
                exp_text = self._apply_rewrites(exp_raw, rewritten_bullets)
                self._add_entries(doc, exp_text)

        proj_parsed = sections.get("projects_parsed")
        proj_raw = sections.get("projects", {}).get("raw", "")
        if proj_parsed or proj_raw:
            self._heading(doc, "PROJECTS")
            if proj_parsed:
                for entry in proj_parsed:
                    self._add_structured_entry(doc, entry, rewritten_bullets)
            else:
                proj_text = self._apply_rewrites(proj_raw, rewritten_bullets)
                self._add_entries(doc, proj_text, trailing_space=False)

        cert_raw = sections.get("certifications", {}).get("raw", "")
        if cert_raw:
            self._add_list_section(doc, "CERTIFICATIONS", cert_raw)

        lang_raw = sections.get("languages", {}).get("raw", "")
        if lang_raw:
            self._add_list_section(doc, "LANGUAGES", lang_raw)

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.getvalue()

    # ---------------------------------------------------------------- #
    # setup
    # ---------------------------------------------------------------- #

    def _setup_page(self, doc):
        for s in doc.sections:
            s.top_margin = Inches(0.4)
            s.bottom_margin = Inches(0.4)
            s.left_margin = Inches(0.45)
            s.right_margin = Inches(0.45)
        self.page_width = Inches(7.6)

    def _setup_base_style(self, doc):
        style: Any = doc.styles["Normal"]
        style.font.name = self.FONT_FAMILY
        style.font.size = Pt(9.5)
        style.font.color.rgb = self.COLOR_BODY_RGB
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.08
        # Force Arial into the style's rFonts slot
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:ascii"), self.FONT_FAMILY)
        rFonts.set(qn("w:hAnsi"), self.FONT_FAMILY)
        rFonts.set(qn("w:eastAsia"), self.FONT_FAMILY)
        rFonts.set(qn("w:cs"), self.FONT_FAMILY)

    def _setup_bullet_style(self, doc):
        existing = [s.name for s in doc.styles]
        if "ResumeBullet" in existing:
            return
        bullet_style = doc.styles.add_style("ResumeBullet", WD_STYLE_TYPE.PARAGRAPH)
        bullet_style.base_style = doc.styles["List Bullet"]
        bullet_style.font.name = self.FONT_FAMILY
        bullet_style.font.size = Pt(9.5)
        bullet_style.font.color.rgb = self.COLOR_BODY_RGB
        rPr = bullet_style.element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:ascii"), self.FONT_FAMILY)
        rFonts.set(qn("w:hAnsi"), self.FONT_FAMILY)
        rFonts.set(qn("w:eastAsia"), self.FONT_FAMILY)
        rFonts.set(qn("w:cs"), self.FONT_FAMILY)
        pf = bullet_style.paragraph_format
        pf.space_after = Pt(1.5)
        pf.line_spacing = 1.05
        pf.left_indent = Inches(0.22)
        pf.first_line_indent = Inches(-0.15)

    def _force_font(self, run, size_pt=None, bold=None, italic=None, color=None):
        """Make Arial actually stick on this run."""
        run.font.name = self.FONT_FAMILY
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:ascii"), self.FONT_FAMILY)
        rFonts.set(qn("w:hAnsi"), self.FONT_FAMILY)
        rFonts.set(qn("w:eastAsia"), self.FONT_FAMILY)
        rFonts.set(qn("w:cs"), self.FONT_FAMILY)
        if size_pt is not None:
            run.font.size = Pt(size_pt)
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic
        if color is not None:
            run.font.color.rgb = color

    # ---------------------------------------------------------------- #
    # header / contact
    # ---------------------------------------------------------------- #

    def _add_header(
        self,
        doc,
        c: dict,
        target_role: str = "",
        resume_title: str = "",
        match_title: bool = False,
    ):
        if c.get("name"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(c["name"].upper())
            self._force_font(r, size_pt=14, bold=True, color=RGBColor(10, 10, 10))

        # A4: never invent a title. Default to the resume's own headline; the
        # job posting's title is used only when the user ticks "match this job"
        # for this export. No title at all beats a fabricated one, so the
        # subtitle line is simply omitted when neither is available.
        job_title = (target_role or "").strip()
        if job_title.lower() == "string":  # legacy placeholder value in stored data
            job_title = ""
        chosen = job_title if (match_title and job_title) else (resume_title or "").strip()
        clean_role = chosen.upper()
        if clean_role:
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sub.paragraph_format.space_after = Pt(3)
            r_sub = p_sub.add_run(clean_role)
            self._force_font(
                r_sub, size_pt=10, bold=True, color=self.COLOR_PRIMARY_RGB
            )

        p_cnt = doc.add_paragraph()
        p_cnt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cnt.paragraph_format.space_after = Pt(6)

        parts = self.contact_parts(c)

        for i, (url, text) in enumerate(parts):
            if i > 0:
                sep = p_cnt.add_run("   |   ")
                self._force_font(sep, size_pt=8.5, color=self.COLOR_MUTED_RGB)
            if url:
                self._add_hyperlink(p_cnt, url, text, size=8.5)
            else:
                run = p_cnt.add_run(text)
                self._force_font(run, size_pt=8.5, color=self.COLOR_BODY_RGB)

    def _add_hyperlink(self, paragraph, url: str, text: str, size: float = 8.5):
        part = paragraph.part
        r_id = part.relate_to(
            url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True
        )
        hyperlink_xml = (
            f'<w:hyperlink {nsdecls("w")} '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            f'r:id="{r_id}">'
            "<w:r><w:rPr>"
            f'<w:rFonts w:ascii="{self.FONT_FAMILY}" w:hAnsi="{self.FONT_FAMILY}" '
            f'w:eastAsia="{self.FONT_FAMILY}" w:cs="{self.FONT_FAMILY}"/>'
            '<w:color w:val="0563C1"/><w:u w:val="single"/>'
            f'<w:sz w:val="{int(size * 2)}"/>'
            "</w:rPr>"
            f'<w:t xml:space="preserve">{self._xml_escape(text)}</w:t>'
            "</w:r></w:hyperlink>"
        )
        paragraph._p.append(parse_xml(hyperlink_xml))

    def _xml_escape(self, text: str) -> str:
        return (
            text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

    # ---------------------------------------------------------------- #
    # section heading
    # ---------------------------------------------------------------- #

    def _heading(self, doc, title: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title.upper())
        self._force_font(r, size_pt=10, bold=True, color=self.COLOR_PRIMARY_RGB)

        pBdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="4" w:space="1" '
            f'w:color="{self.COLOR_PRIMARY_HEX}"/></w:pBdr>'
        )
        p._p.get_or_add_pPr().append(pBdr)

    # ---------------------------------------------------------------- #
    # flowing text (summary)
    # ---------------------------------------------------------------- #

    def _add_flow_text(self, doc, text: str):
        joined = self.summary_text(text)
        if not joined:
            return
        p = doc.add_paragraph(joined)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.08
        for run in p.runs:
            self._force_font(run, size_pt=9.5, color=self.COLOR_BODY_RGB)

    # ---------------------------------------------------------------- #
    # education (structured)
    # ---------------------------------------------------------------- #

    def _add_education_entry(self, doc, entry: dict):
        """
        Renders one education_parsed item:
          {dates, school, degree, details[]}
        Layout:
          School                                         2022 - 2026
          Degree - honors
          extra detail lines
        """
        parts = self.education_parts(entry)
        primary = parts["primary"]
        secondary = parts["secondary"]
        dates = parts["dates"]
        details = parts["details"]
        if not primary:
            return

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        if dates:
            p.paragraph_format.tab_stops.add_tab_stop(
                self.page_width, WD_TAB_ALIGNMENT.RIGHT
            )

        r = p.add_run(primary)
        self._force_font(r, size_pt=9.5, bold=True, color=self.COLOR_BODY_RGB)

        if dates:
            p.add_run("\t")
            r_d = p.add_run(dates)
            self._force_font(r_d, size_pt=9, color=self.COLOR_MUTED_RGB)

        if secondary:
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(1)
            r2 = p2.add_run(secondary)
            self._force_font(r2, size_pt=9, italic=True, color=self.COLOR_MUTED_RGB)

        for d in details:
            d = (d or "").strip()
            if not d or self._is_layout_label(d):
                continue
            pd = doc.add_paragraph()
            pd.paragraph_format.space_after = Pt(0)
            rd = pd.add_run(d)
            self._force_font(rd, size_pt=9, color=self.COLOR_BODY_RGB)

        if doc.paragraphs:
            doc.paragraphs[-1].paragraph_format.space_after = Pt(2)

    # ---------------------------------------------------------------- #
    # structured entry (experience / projects)
    # ---------------------------------------------------------------- #

    def _add_structured_entry(self, doc, entry: dict, rewritten: dict | None):
        parts = self.entry_parts(entry)
        title = parts["title"]
        subtitle = parts["subtitle"]
        dates = parts["dates"]
        bullets = entry.get("bullets") or []

        if not title and not bullets:
            return

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        if dates:
            p.paragraph_format.tab_stops.add_tab_stop(
                self.page_width, WD_TAB_ALIGNMENT.RIGHT
            )

        r_title = p.add_run(title)
        self._force_font(r_title, size_pt=9.5, bold=True, color=self.COLOR_BODY_RGB)

        if dates:
            p.add_run("\t")
            r_date = p.add_run(dates)
            self._force_font(r_date, size_pt=9, color=self.COLOR_MUTED_RGB)

        if subtitle:
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(1)
            r_sub = p2.add_run(subtitle)
            self._force_font(
                r_sub, size_pt=9, italic=True, color=self.COLOR_MUTED_RGB
            )

        for text in self.entry_bullets(bullets, rewritten):
            bp = doc.add_paragraph(text, style="ResumeBullet")
            for run in bp.runs:
                self._force_font(run, size_pt=9.5, color=self.COLOR_BODY_RGB)

        if doc.paragraphs:
            doc.paragraphs[-1].paragraph_format.space_after = Pt(3)

    # ---------------------------------------------------------------- #
    # safe line-based fallback
    # ---------------------------------------------------------------- #

    def _add_entries(self, doc, text: str, trailing_space: bool = True):
        lines = self._clean_lines(text)
        for line in lines:
            if self.BULLET_RE.match(line):
                content = self.BULLET_RE.sub("", line).strip()
                bp = doc.add_paragraph(content, style="ResumeBullet")
                for run in bp.runs:
                    self._force_font(run, size_pt=9.5, color=self.COLOR_BODY_RGB)
                continue

            is_new_entry = bool(self.DATE_RE.match(line))
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4) if is_new_entry else Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(line)
            self._force_font(
                r,
                size_pt=9.5,
                bold=is_new_entry,
                color=self.COLOR_BODY_RGB,
            )

        if trailing_space and doc.paragraphs:
            doc.paragraphs[-1].paragraph_format.space_after = Pt(4)

    def _add_list_section(self, doc, title: str, text: str):
        """Render a simple section (certifications, languages) line by line."""
        lines = self._clean_lines(text)
        if not lines:
            return
        self._heading(doc, title)
        self._add_entries(doc, "\n".join(lines), trailing_space=False)

    # ---------------------------------------------------------------- #
    # skills
    # ---------------------------------------------------------------- #

    def _add_skills(self, doc, text: str, missing_keywords: list[str] | None):
        # A3: only user-confirmed technical skills reach the export. They fold
        # into the matching labeled line ("Databases: MySQL, SQL"); anything
        # without a matching line falls back to "Additional Skills". Soft
        # skills never get here (see _filter_keywords).
        lines, leftover = self.skills_lines(text, missing_keywords)
        for line in lines:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.08
            if ":" in line:
                label, _, rest = line.partition(":")
                r_label = p.add_run(label.strip() + ": ")
                self._force_font(
                    r_label, size_pt=9.5, bold=True, color=self.COLOR_BODY_RGB
                )
                r_rest = p.add_run(rest.strip())
                self._force_font(r_rest, size_pt=9.5, color=self.COLOR_BODY_RGB)
            else:
                r = p.add_run(line)
                self._force_font(r, size_pt=9.5, color=self.COLOR_BODY_RGB)

        if leftover:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.08
            r_label = p.add_run("Additional Skills: ")
            self._force_font(
                r_label, size_pt=9.5, bold=True, color=self.COLOR_BODY_RGB
            )
            r_rest = p.add_run(", ".join(leftover))
            self._force_font(r_rest, size_pt=9.5, color=self.COLOR_BODY_RGB)

    # Skills, rewrite, camel-case, and line-normalization helpers live in
    # ResumeContentBuilder (app/export/content.py) so the PDF exporter applies
    # exactly the same content rules.
