import io
import unittest

from docx import Document
from docx.oxml.ns import qn

from app.export.docx_exporter import DocxExporter


def document_text(data: bytes) -> str:
    """Return every paragraph's visible text, including hyperlink runs.

    python-docx's ``Paragraph.text`` skips text nested inside ``w:hyperlink``,
    so read the XML text nodes instead.
    """
    doc = Document(io.BytesIO(data))
    lines = []
    for paragraph in doc.paragraphs:
        lines.append("".join(node.text or "" for node in paragraph._p.iter(qn("w:t"))))
    return "\n".join(lines)


def make_sections() -> dict:
    return {
        "contact_parsed": {
            "name": "Jane Doe",
            "email": "jane@example.com",
            "phone": "555-123-4567",
            "location": "Manila, PH",
            "portfolio": "janedoe.dev",
            "linkedin": "linkedin.com/in/janedoe",
            "github": "github.com/janedoe",
        },
        "summary": {"raw": "Backend developer."},
        "skills": {"raw": "Languages: Python, JavaScript, TypeScript, SQL"},
        "education_parsed": [],
        "experience_parsed": [
            {
                "title": "Backend Engineer",
                "company": "Acme",
                "dates": "2023 - 2025",
                "bullets": [
                    "Built APIs with JavaScript, TypeScript, PostgreSQL, MongoDB, and GitHub Actions",
                    "Used FastAPI and GraphQL with SQLAlchemy and SQLModel on AWS",
                    "Shipped iOS and Android clients with WebSockets and NoSQL stores",
                    "Practiced DevOps and improved a WordWord pipeline",
                ],
            }
        ],
        "projects_parsed": [],
        "certifications": {
            "raw": (
                "AWS Certified Cloud Practitioner - 2024\n"
                "Certified Kubernetes Administrator"
            )
        },
        "languages": {"raw": "English - Fluent\nFilipino - Native"},
    }


class DocxExporterCamelCaseTests(unittest.TestCase):
    def setUp(self):
        self.text = document_text(DocxExporter().export_tailored(make_sections()))

    def test_canonical_tech_names_are_not_split(self):
        for expected in (
            "JavaScript",
            "TypeScript",
            "PostgreSQL",
            "MongoDB",
            "GitHub",
            "GitHub Actions",
            "FastAPI",
            "GraphQL",
            "SQLAlchemy",
            "SQLModel",
            "iOS",
            "WebSockets",
            "NoSQL",
            "DevOps",
        ):
            with self.subTest(term=expected):
                self.assertIn(expected, self.text)

    def test_broken_camel_case_spellings_are_absent(self):
        for broken in (
            "Java Script",
            "Type Script",
            "Postgre SQL",
            "Mongo DB",
            "Git Hub",
            "Fast API",
            "Graph QL",
            "i OS",
            "Web Sockets",
            "No SQL",
            "Dev Ops",
        ):
            with self.subTest(broken=broken):
                self.assertNotIn(broken, self.text)

    def test_unknown_camel_case_still_splits(self):
        self.assertIn("Word Word", self.text)
        self.assertNotIn("WordWord", self.text)


class DocxExporterDroppedContentTests(unittest.TestCase):
    def test_linkedin_and_github_both_render_in_header(self):
        text = document_text(DocxExporter().export_tailored(make_sections()))

        self.assertIn("linkedin.com/in/janedoe", text)
        self.assertIn("github.com/janedoe", text)

    def test_certifications_and_languages_render(self):
        text = document_text(DocxExporter().export_tailored(make_sections()))

        self.assertIn("CERTIFICATIONS", text)
        self.assertIn("AWS Certified Cloud Practitioner - 2024", text)
        self.assertIn("Certified Kubernetes Administrator", text)
        self.assertIn("LANGUAGES", text)
        self.assertIn("English - Fluent", text)
        self.assertIn("Filipino - Native", text)

    def test_empty_sections_do_not_add_headings(self):
        sections = make_sections()
        sections.pop("certifications")
        sections.pop("languages")

        text = document_text(DocxExporter().export_tailored(sections))

        self.assertNotIn("CERTIFICATIONS", text)
        self.assertNotIn("LANGUAGES", text)


class DocxExporterGlyphTests(unittest.TestCase):
    """Real-resume click-through: Groq rewrites contained U+2011 non-breaking
    hyphens, which render as tofu boxes (□) in DOCX viewers."""

    def _section_with_ai_bullet(self) -> dict:
        sections = make_sections()
        sections["projects_parsed"] = [
            {
                "title": "Online Examination System",
                "company": "",
                "dates": "2025 - 2026",
                "bullets": [
                    "Led end-to-end development of a secure exam platform using React.js and NestJS",
                ],
            }
        ]
        return sections

    def test_ai_output_nonbreaking_hyphens_are_normalized(self):
        rewritten = {
            "Led end-to-end development of a secure exam platform using React.js and NestJS": {
                "ai_output": "Led end\u2011to\u2011end development of a web\u2011based exam platform with NestJS",
            }
        }
        text = document_text(
            DocxExporter().export_tailored(self._section_with_ai_bullet(), rewritten_bullets=rewritten)
        )
        self.assertIn("Led end-to-end development of a web-based exam platform with NestJS", text)
        self.assertNotIn("\u2011", text)

    def test_raw_bullet_nonbreaking_hyphens_are_normalized(self):
        sections = self._section_with_ai_bullet()
        sections["projects_parsed"][0]["bullets"] = [
            "Built an arcade\u2011style game with real\u2011time scoring on problem\u2011solving"
        ]
        text = document_text(DocxExporter().export_tailored(sections))
        self.assertIn("arcade-style game with real-time scoring on problem-solving", text)

    def test_en_and_em_dashes_are_preserved(self):
        # The parser's canonical separator is " — " — it must not be clobbered.
        sections = make_sections()
        sections["education_parsed"] = [
            {
                "dates": "2021 - 2022",
                "school": "State University",
                "degree": "Senior High School, ABM Strand — Graduated with Highest Honors",
                "details": [],
            }
        ]
        text = document_text(DocxExporter().export_tailored(sections))
        self.assertIn("ABM Strand — Graduated with Highest Honors", text)

    def test_nestjs_is_not_split(self):
        sections = self._section_with_ai_bullet()
        text = document_text(DocxExporter().export_tailored(sections))
        self.assertIn("NestJS", text)
        self.assertNotIn("Nest JS", text)


class KeywordFilterTests(unittest.TestCase):
    """A3: confirmed missing keywords reach the exported Skills section."""

    def test_filter_drops_fluff_duplicates_and_known_skills(self):
        exporter = DocxExporter()
        skills = "Languages: Python, JavaScript, TypeScript, SQL"
        out = exporter._filter_keywords(
            skills,
            ["Docker", "docker", "SQL", "great team player", "Go", "x" * 30],
        )
        # "docker" deduped, "SQL" already in skills, fluff and length
        # bounds dropped, "Go" too short
        self.assertEqual(out, ["Docker"])

    def test_filter_preserves_canonical_spelling(self):
        exporter = DocxExporter()
        out = exporter._filter_keywords("", ["postgresql", "javascript", "Kubernetes"])
        self.assertIn("PostgreSQL", out)
        self.assertIn("JavaScript", out)
        self.assertIn("Kubernetes", out)
        self.assertNotIn("Sql", out)
        self.assertNotIn("Postgresql", out)

    def test_filter_does_not_treat_mysql_as_sql(self):
        # Live regression: "SQL" was dropped because "MySQL" contains "sql"
        # as a substring. Whole-word matching keeps the confirmed skill.
        out = DocxExporter()._filter_keywords(
            "Databases: MySQL, PostgreSQL", ["SQL", "Java"]
        )
        self.assertEqual(out, ["SQL", "Java"])

    def test_export_appends_additional_skills_line(self):
        sections = make_sections()
        text = document_text(
            DocxExporter().export_tailored(
                sections, missing_keywords=["Docker", "Kubernetes"]
            )
        )
        self.assertIn("Additional Skills: Docker, Kubernetes", text)

    def test_export_skips_keyword_already_in_skills(self):
        sections = make_sections()
        text = document_text(
            DocxExporter().export_tailored(
                sections, missing_keywords=["SQL", "Docker"]
            )
        )
        # SQL is already in the skills line -- not duplicated
        self.assertIn("Additional Skills: Docker", text)
        self.assertNotIn("Additional Skills: SQL", text)
        self.assertNotIn("Additional Skills: Docker, SQL", text)

    def test_export_without_keywords_has_no_additional_skills_line(self):
        text = document_text(DocxExporter().export_tailored(make_sections()))
        self.assertNotIn("Additional Skills", text)

    def test_hard_skill_folds_into_matching_category(self):
        sections = make_sections()
        sections["skills"] = {"raw": (
            "Languages: Python, JavaScript\n"
            "Databases: MySQL, PostgreSQL\n"
            "DevOps & Tools: Docker, Git"
        )}
        text = document_text(
            DocxExporter().export_tailored(
                sections, missing_keywords=["SQL", "Linux"]
            )
        )
        # folded into the existing labeled lines, not bolted on
        self.assertIn("Databases: MySQL, PostgreSQL, SQL", text)
        self.assertIn("DevOps & Tools: Docker, Git, Linux", text)
        self.assertNotIn("Additional Skills", text)

    def test_unmatched_hard_skill_falls_back_to_additional_skills(self):
        sections = make_sections()
        sections["skills"] = {"raw": "Languages: Python"}
        text = document_text(
            DocxExporter().export_tailored(
                sections, missing_keywords=["Kubernetes"]
            )
        )
        self.assertIn("Additional Skills: Kubernetes", text)

    def test_soft_skills_are_never_added_to_skills(self):
        # "Innovative, Interpersonal" belongs in bullets, not a skills list.
        text = document_text(
            DocxExporter().export_tailored(
                make_sections(),
                missing_keywords=["Communication", "Interpersonal"],
            )
        )
        self.assertNotIn("Additional Skills", text)
        self.assertNotIn("Interpersonal", text)

    def test_filter_excludes_soft_skills(self):
        out = DocxExporter()._filter_keywords("", ["Communication", "Docker"])
        self.assertEqual(out, ["Docker"])

    def test_paradigm_is_not_filed_under_a_tools_line(self):
        # Live bug: confirmed OOP rendered as
        # "DevOps & Tools: Docker, Nginx, VPS, Git, GitHub, Figma,
        #  Object-Oriented Programming" — a paradigm is not a DevOps tool.
        sections = make_sections()
        sections["skills"] = {"raw": (
            "Languages: Python, JavaScript\n"
            "DevOps & Tools: Docker, Git, Figma"
        )}
        text = document_text(
            DocxExporter().export_tailored(
                sections, missing_keywords=["Object-Oriented Programming"]
            )
        )
        self.assertNotIn(
            "DevOps & Tools: Docker, Git, Figma, Object-Oriented Programming", text
        )
        self.assertIn("Additional Skills: Object-Oriented Programming", text)

    def test_real_tools_still_fold_into_a_tools_line(self):
        sections = make_sections()
        sections["skills"] = {"raw": "DevOps & Tools: Docker, Git"}
        text = document_text(
            DocxExporter().export_tailored(sections, missing_keywords=["Figma"])
        )
        self.assertIn("DevOps & Tools: Docker, Git, Figma", text)
        self.assertNotIn("Additional Skills", text)


class HeaderTitleTests(unittest.TestCase):
    """A4: the export shows the resume's own title, never the job's by default."""

    def _sections_with_title(self, title: str = "Full Stack Developer") -> dict:
        sections = make_sections()
        sections["contact_parsed"]["title"] = title
        return sections

    def test_own_title_is_kept_when_a_job_is_attached(self):
        text = document_text(
            DocxExporter().export_tailored(
                self._sections_with_title(), target_role="Frontend Apprentice"
            )
        )
        self.assertIn("FULL STACK DEVELOPER", text)
        self.assertNotIn("FRONTEND APPRENTICE", text)

    def test_job_title_is_used_only_when_explicitly_requested(self):
        text = document_text(
            DocxExporter().export_tailored(
                self._sections_with_title(),
                target_role="Frontend Apprentice",
                match_title=True,
            )
        )
        self.assertIn("FRONTEND APPRENTICE", text)
        self.assertNotIn("FULL STACK DEVELOPER", text)

    def test_match_title_without_a_job_keeps_the_own_title(self):
        text = document_text(
            DocxExporter().export_tailored(
                self._sections_with_title(), target_role="", match_title=True
            )
        )
        self.assertIn("FULL STACK DEVELOPER", text)

    def test_no_title_is_never_fabricated(self):
        # The old hardcoded "FULL STACK DEVELOPER" fallback printed a title
        # the candidate never claimed, for every user without a job attached.
        sections = make_sections()
        text = document_text(DocxExporter().export_tailored(sections))
        self.assertNotIn("FULL STACK DEVELOPER", text)
        self.assertIn("JANE DOE", text)  # the rest of the header still renders

    def test_legacy_resume_recovers_title_from_contact_raw(self):
        # Resumes parsed before A4 have no contact_parsed["title"]; the title
        # must still be recovered from the stored header text on export.
        sections = make_sections()
        sections["contact_parsed"].pop("title", None)
        sections["contact"] = {
            "raw": (
                "Jane Doe\nFULL STACK DEVELOPER\n"
                "jane@example.com | 555-123-4567"
            )
        }
        text = document_text(
            DocxExporter().export_tailored(sections, target_role="Data Analyst")
        )
        self.assertIn("FULL STACK DEVELOPER", text)
        self.assertNotIn("DATA ANALYST", text)

    def test_summary_keeps_the_word_aspiring(self):
        sections = make_sections()
        sections["summary"] = {"raw": "Aspiring Full Stack Developer skilled in React.js."}
        text = document_text(DocxExporter().export_tailored(sections))
        self.assertIn("Aspiring Full Stack Developer", text)

    def test_layout_labels_are_configurable(self):
        class CustomLabels(DocxExporter):
            LAYOUT_LABELS = DocxExporter.LAYOUT_LABELS | {"leftover"}

        sections = make_sections()
        sections["skills"] = {"raw": "Leftover\nLanguages: Python"}
        default_text = document_text(DocxExporter().export_tailored(sections))
        custom_text = document_text(CustomLabels().export_tailored(sections))
        self.assertIn("Leftover", default_text)
        self.assertNotIn("Leftover", custom_text)


if __name__ == "__main__":
    unittest.main()
