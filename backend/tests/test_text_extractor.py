import base64
import io
import unittest

import docx
import pymupdf as fitz

from app.parser.text_extractor import TextExtractor


WORDS = [
    "alpha", "bravo", "charlie", "delta", "echo", "foxtrot", "golf", "hotel",
    "india", "juliet", "kilo", "lima", "mike", "november", "oscar", "papa",
]

# 1x1 transparent PNG used for the DOCX image test.
PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def make_pdf(pages=1, columns=False, image=False, header_line=None, footer_line=None):
    doc = fitz.open()
    for page_index in range(pages):
        page = doc.new_page(width=612, height=792)
        y = 72
        if header_line:
            page.insert_text((50, y), header_line)
            y += 24

        if columns:
            for offset, word in enumerate(WORDS):
                page.insert_text((50, y + offset * 18), word)
                page.insert_text((350, y + offset * 18), word)
        else:
            for offset in range(40):
                page.insert_text((50, y + offset * 16), WORDS[offset % len(WORDS)])

        if image:
            pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 8, 8))
            pix.clear_with(255)
            page.insert_image(fitz.Rect(500, 700, 540, 740), pixmap=pix)

        if footer_line:
            page.insert_text((50, 760), footer_line.format(page=page_index + 1))

    data = doc.tobytes()
    doc.close()
    return data


def make_docx(header=None, image=False, fonts=(), icon_font=False):
    document = docx.Document()
    if header:
        document.sections[0].header.paragraphs[0].text = header
    for font in fonts:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"Text set in {font}")
        run.font.name = font
    if icon_font:
        paragraph = document.add_paragraph()
        run = paragraph.add_run("\uf2b9")
        run.font.name = "Font Awesome 5 Free"
    if image:
        document.add_picture(io.BytesIO(PNG_1PX))
    document.add_paragraph("Built reliable services in Python.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class PdfMetadataTests(unittest.TestCase):
    """B9: PDFs now expose layout metadata instead of always passing."""

    extractor = TextExtractor()

    def test_single_column_pdf_is_not_flagged(self):
        result = self.extractor.extract(make_pdf(), ".pdf")
        meta = result["metadata"]
        self.assertEqual(meta["page_count"], 1)
        self.assertFalse(meta["has_multi_column"])
        self.assertFalse(meta["has_images"])
        self.assertEqual(meta["image_count"], 0)
        self.assertGreaterEqual(meta["font_count"], 1)
        self.assertFalse(meta["has_headers"])
        self.assertFalse(meta["has_footers"])
        self.assertIn("alpha", result["raw_text"])

    def test_two_column_pdf_is_detected(self):
        result = self.extractor.extract(make_pdf(columns=True), ".pdf")
        self.assertTrue(result["metadata"]["has_multi_column"])

    def test_images_are_counted(self):
        result = self.extractor.extract(make_pdf(image=True), ".pdf")
        meta = result["metadata"]
        self.assertTrue(meta["has_images"])
        self.assertGreaterEqual(meta["image_count"], 1)

    def test_repeated_contact_header_is_detected(self):
        data = make_pdf(pages=2, header_line="alex.doe@example.com")
        meta = self.extractor.extract(data, ".pdf")["metadata"]
        self.assertTrue(meta["has_headers"])
        self.assertFalse(meta["has_footers"])

    def test_page_numbers_are_not_headers_or_footers(self):
        data = make_pdf(pages=2, footer_line="Page {page}")
        meta = self.extractor.extract(data, ".pdf")["metadata"]
        self.assertFalse(meta["has_headers"])
        self.assertFalse(meta["has_footers"])


class DocxMetadataTests(unittest.TestCase):
    """B9: DOCX images, fonts, and headers are visible to the checker."""

    extractor = TextExtractor()

    def test_plain_docx_is_clean(self):
        meta = self.extractor.extract(make_docx(), ".docx")["metadata"]
        self.assertFalse(meta["has_headers"])
        self.assertFalse(meta["has_images"])
        self.assertEqual(meta["image_count"], 0)
        self.assertEqual(meta["font_count"], 0)
        self.assertFalse(meta["has_icon_font"])
        self.assertFalse(meta["has_multi_column"])

    def test_header_is_detected(self):
        meta = self.extractor.extract(make_docx(header="alex.doe@example.com"), ".docx")["metadata"]
        self.assertTrue(meta["has_headers"])

    def test_image_is_counted(self):
        meta = self.extractor.extract(make_docx(image=True), ".docx")["metadata"]
        self.assertTrue(meta["has_images"])
        self.assertEqual(meta["image_count"], 1)

    def test_font_variety_is_counted(self):
        meta = self.extractor.extract(make_docx(fonts=("Arial", "Calibri", "Georgia")), ".docx")["metadata"]
        self.assertEqual(meta["font_count"], 3)

    def test_icon_font_is_detected(self):
        meta = self.extractor.extract(make_docx(icon_font=True), ".docx")["metadata"]
        self.assertTrue(meta["has_icon_font"])


if __name__ == "__main__":
    unittest.main()
